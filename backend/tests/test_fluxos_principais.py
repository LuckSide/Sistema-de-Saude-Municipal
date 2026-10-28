"""Suíte de Testes Automatizados de Integração dos Fluxos Principais do Sistema:
1. Autenticação e Permissões (RBAC)
2. Listagem e Busca de Pacientes e Profissionais (Filtro Médicos vs. Motoristas)
3. Inclusão de Exames na Fila de Espera (Justificativa >= 12 chars, Prioridade e SIGTAP)
4. Rotas de Agendamentos e Visão Semanal
5. Registro de Auditoria e Logs
"""

from datetime import date, datetime, timedelta, timezone
import json
import pytest
from fastapi import HTTPException
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.security import hash_password, verify_password, create_access_token
from app.services.permissoes import tem_tela, tem_acao, need_perm, need_any_perm
from app.models.user import User
from app.models.cadastros_saude import Medico, UnidadeSaude
from app.models.tfd import Motorista, Paciente, ExameFila, Viagem, ViagemPaciente
from app.models.lab_exames import SigtapProcedimento
from app.models.lab_fase6 import FilaPsf
from app.models.audit import AuditLog
from app.schemas.tfd import MotoristaIn, PacienteIn, ExameIn
from app.schemas.cadastros_saude import MedicoIn
from app.api.v1.tfd import (
    listar_motoristas,
    criar_motorista,
    atualizar_motorista,
    listar_pacientes,
    criar_paciente,
    criar_exame,
    atualizar_exame,
    listar_exames,
    resolver_medico_unidade,
)
from app.api.v1.cadastros_saude import listar_medicos, _validar_nao_motorista
from app.core.audit import parse_device_from_ua, registrar_log_pesquisa, registrar_log_exportacao
from app.services.log_rotacao import obter_estatisticas_auditoria, executar_rotacao_logs


@pytest.fixture
def db():
    test_engine = create_engine("sqlite:///:memory:")
    SQLModel.metadata.create_all(test_engine)
    with Session(test_engine) as session:
        yield session


# =====================================================================
# 1. AUTENTICAÇÃO E PERMISSÕES (RBAC)
# =====================================================================

def test_auth_rbac_hashing_and_jwt():
    senha_pura = "SenhaSegura@2026"
    hash_gerado = hash_password(senha_pura)
    assert hash_gerado != senha_pura
    assert verify_password(senha_pura, hash_gerado) is True
    assert verify_password("SenhaErrada", hash_gerado) is False

    token = create_access_token("admin_user")
    assert isinstance(token, str) and len(token) > 20


def test_auth_rbac_permissoes_hierarquicas():
    # Superuser tem acesso irrestrito
    super_user = User(username="superadmin", is_superuser=True)
    assert tem_tela(super_user, "qualquer_tela") is True
    assert tem_acao(super_user, "qualquer_tela", "qualquer_acao") is True

    # Usuário com módulo pai "tfd" herda telas filhas regulares
    user_tfd = User(
        username="operador_tfd",
        is_superuser=False,
        telas="tfd",
        acoes='{"tfd": ["ver", "criar", "editar"]}',
    )
    assert tem_tela(user_tfd, "tfd") is True
    assert tem_tela(user_tfd, "tfd_frota") is True
    assert tem_tela(user_tfd, "tfd_destinos") is True

    # Telas EXCLUSIVAS (ex: tfd_viagens_passadas, apac_regulacao) NÃO são herdadas automaticamente
    assert tem_tela(user_tfd, "tfd_viagens_passadas") is False
    assert tem_tela(user_tfd, "apac_regulacao") is False

    # Concessão explícita de tela exclusiva
    user_especial = User(
        username="auditor",
        is_superuser=False,
        telas="tfd,tfd_viagens_passadas",
        acoes='{"tfd_viagens_passadas": ["ver", "editar"]}',
    )
    assert tem_tela(user_especial, "tfd_viagens_passadas") is True
    assert tem_acao(user_especial, "tfd_viagens_passadas", "editar") is True


def test_auth_rbac_guards():
    check_frota = need_any_perm("tfd_frota", "tfd_frota_gestao")
    user_sem_acesso = User(username="sem_acesso", is_superuser=False, telas="")

    with pytest.raises(HTTPException) as exc:
        check_frota(user_sem_acesso)
    assert exc.value.status_code == 403

    user_com_acesso = User(
        username="gestor_frota",
        is_superuser=False,
        telas="tfd_frota",
    )
    res = check_frota(user_com_acesso)
    assert res.username == "gestor_frota"


# =====================================================================
# 2. LISTAGEM E BUSCA: MOTORISTAS VS. MÉDICOS E PACIENTES
# =====================================================================

def test_motoristas_sem_erro_500_mesmo_sem_cnh_completa(db):
    # Insere motorista sem CNH (como os importados do PEC ou higienizados)
    mot = Motorista(
        nome="SEBASTIAO ANTONIO PEREIRA",
        cpf="123.456.789-00",
        cnh_numero="",
        cnh_categoria="",
        cnh_validade=None,
        telefone="",
        ativo=True,
    )
    db.add(mot)
    db.commit()

    # Chama listar_motoristas - deve retornar HTTP 200 / Pagina[MotoristaOut] sem ResponseValidationError
    res = listar_motoristas(q="", page=1, limit=100, db=db)
    assert res["total"] == 1
    item = res["data"][0]
    assert item.nome == "SEBASTIAO ANTONIO PEREIRA"
    assert item.cpf == "123.456.789-00"
    assert item.cnh_numero == ""
    assert item.cnh_validade is None


def test_filtro_medicos_exclui_motoristas(db):
    # 1. Médico clínico real
    med_clinico = Medico(
        nome="Dr. Roberto Campos",
        crm="CRM/MG 45890",
        cbo_codigo="225125",
        cbo_descricao="Médico Clínico Geral",
        ativo=True,
    )
    # 2. Profissional que é motorista (CBO 7823)
    med_motorista = Medico(
        nome="Joaquim do Transporte",
        crm="",
        cbo_codigo="782310",
        cbo_descricao="Motorista de Ambulância",
        ativo=True,
    )
    db.add(med_clinico)
    db.add(med_motorista)
    db.commit()

    # Validação do backend rejeita cadastrar médico motorista
    with pytest.raises(HTTPException) as exc:
        _validar_nao_motorista(MedicoIn(nome="José Motorista", cbo_codigo="782320"))
    assert exc.value.status_code == 400

    # listar_medicos NÃO deve retornar o motorista
    lista = listar_medicos(q="", somente_ativos=True, db=db)
    nomes = [m.nome for m in lista]
    assert "Dr. Roberto Campos" in nomes
    assert "Joaquim do Transporte" not in nomes


def test_busca_pacientes_ordenacao_inteligente(db):
    pac_inicio = Paciente(nome="Mateus da Silva", cpf="111.111.111-11", municipio="Guapé")
    pac_meio = Paciente(nome="João Mateus Pereira", cpf="222.222.222-22", municipio="Guapé")
    db.add(pac_inicio)
    db.add(pac_meio)
    db.commit()

    res = listar_pacientes(q="Mateus", page=1, limit=10, db=db)
    assert res["total"] == 2
    # Quem começa com o termo deve aparecer em primeiro
    assert res["data"][0].nome == "Mateus da Silva"
    assert res["data"][1].nome == "João Mateus Pereira"


# =====================================================================
# 3. INCLUSÃO DE EXAMES NA FILA DE ESPERA (JUSTIFICATIVA >= 12 CHARS, P1 E SIGTAP)
# =====================================================================

def test_validacao_justificativa_clinica_obrigatoria_minimo_12_caracteres():
    # Justificativa menor que 12 caracteres deve falhar com ValueError
    with pytest.raises(ValueError, match="Justificativa clínica obrigatória"):
        ExameIn(
            paciente_id=1,
            paciente_nome="Maria Teste",
            exame="HEMOGRAMA COMPLETO",
            prioridade="urgente",
            justificativa="Curto",
            data_pedido=date.today(),
        )

    # Justificativa com 12 ou mais caracteres passa normalmente
    exame_ok = ExameIn(
        paciente_id=1,
        paciente_nome="Maria Teste",
        exame="HEMOGRAMA COMPLETO",
        prioridade="urgente",
        justificativa="Paciente com suspeita de anemia grave",
        data_pedido=date.today(),
        cid_codigo="R10.4",
        cid_descricao="Outras dores abdominais e as não especificadas",
    )
    assert exame_ok.justificativa == "Paciente com suspeita de anemia grave"
    assert exame_ok.cid_codigo == "R10.4"


def test_inclusao_exame_com_id_imutavel_e_prioridade(db):
    pac = Paciente(nome="Carlos de Paula", cpf="333.333.333-33", municipio="Guapé")
    db.add(pac)
    db.commit()

    ex_in = ExameIn(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="0202010180 · DOSAGEM DE GLICOSE",
        codigo_procedimento="0202010180",
        tipo_fila="Exames Laboratoriais",
        prioridade="urgente",
        justificativa="Glicemia em jejum alterada para monitoramento",
        cid_codigo="E11",
        cid_descricao="Diabetes mellitus não-insulino-dependente",
        data_pedido=date.today(),
    )
    ator = User(username="operador", is_superuser=True)
    exame_salvo = criar_exame(ex_in, db=db, ator=ator)

    assert exame_salvo.id is not None
    assert exame_salvo.codigo_curto is not None
    assert exame_salvo.prioridade == "urgente"
    assert exame_salvo.status == "aguardando"
    assert exame_salvo.codigo_procedimento == "0202010180"
    assert exame_salvo.cid_codigo == "E11"
    assert exame_salvo.cid_descricao == "Diabetes mellitus não-insulino-dependente"


def test_busca_cids_tfd(db):
    from app.models.lab_exames import SigtapCid
    from app.api.v1.tfd import buscar_cids_tfd

    # Inserir alguns CIDs de teste
    c1 = SigtapCid(procedimento_codigo="0202010180", cid="E11.9", descricao="Diabetes mellitus sem complicações")
    c2 = SigtapCid(procedimento_codigo="0202010180", cid="I10", descricao="Hipertensão essencial")
    db.add(c1)
    db.add(c2)
    db.commit()

    ator = User(username="operador", is_superuser=True)
    res = buscar_cids_tfd(q="Diabetes", db=db, ator=ator)
    assert len(res) >= 1
    assert any(c["codigo"] == "E11.9" for c in res)

    res_codigo = buscar_cids_tfd(q="I10", db=db, ator=ator)
    assert len(res_codigo) >= 1
    assert any(c["codigo"] == "I10" for c in res_codigo)


# =====================================================================
# 4. ROTAS DE AGENDAMENTOS E VISÃO SEMANAL
# =====================================================================

def test_agendamento_de_exame_com_medico_e_unidade(db):
    pac = Paciente(nome="Ana Carolina", cpf="444.444.444-44", municipio="Guapé")
    uni = UnidadeSaude(nome="PSF Bela Vista", cnes="7654321", ativo=True)
    med = Medico(nome="Dra. Beatriz Santos", crm="CRM/MG 55443", cbo_codigo="225125", ativo=True)
    db.add(pac)
    db.add(uni)
    db.add(med)
    db.commit()

    # Cria pedido na fila
    ex_in = ExameIn(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="ELETROCARDIOGRAMA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        justificativa="Avaliação pré-operatória de rotina",
        data_pedido=date.today(),
    )
    ator = User(username="operador", is_superuser=True)
    exame = criar_exame(ex_in, db=db, ator=ator)

    # Agenda para uma data
    hoje = date.today()
    ex_agenda_in = ExameIn(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="ELETROCARDIOGRAMA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        justificativa="Avaliação pré-operatória de rotina",
        data_pedido=date.today(),
        data_agendada=hoje,
        hora_agendada="08:30",
        medico_id=med.id,
        unidade_saude_id=uni.id,
        destino="Guapé",
    )
    ex_agendado = atualizar_exame(
        exame.id,
        data=ex_agenda_in,
        db=db,
        ator=ator,
    )

    assert ex_agendado.status == "agendado"
    assert ex_agendado.data_agendada == hoje
    assert ex_agendado.hora_agendada == "08:30"
    assert ex_agendado.medico_id == med.id
    assert ex_agendado.unidade_saude_id == uni.id


def test_bloqueio_viagem_retroativa_sem_permissao():
    user_sem_perm = User(
        username="operador_comum",
        is_superuser=False,
        permite_retroativo=True,
        telas="tfd",
    )
    # Não possui a permissão exclusiva "tfd_viagens_passadas"
    pode_editar = (
        user_sem_perm.is_superuser
        or (
            user_sem_perm.permite_retroativo is not False
            and tem_tela(user_sem_perm, "tfd_viagens_passadas")
        )
    )
    assert pode_editar is False

    user_com_perm = User(
        username="supervisora",
        is_superuser=False,
        permite_retroativo=True,
        telas="tfd,tfd_viagens_passadas",
    )
    pode_editar_super = (
        user_com_perm.is_superuser
        or (
            user_com_perm.permite_retroativo is not False
            and tem_tela(user_com_perm, "tfd_viagens_passadas")
        )
    )
    assert pode_editar_super is True


# =====================================================================
# 5. REGISTRO DE AUDITORIA E LOGS
# =====================================================================

def test_auditoria_device_parsing_accuracy():
    ua_pc = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0 Safari/537.36"
    assert parse_device_from_ua(ua_pc) == "Chrome no Windows (Desktop)"

    ua_celular = "Mozilla/5.0 (Linux; Android 14; Pixel 8) Chrome/120.0.0.0 Mobile Safari/537.36"
    assert parse_device_from_ua(ua_celular) == "Chrome no Android (Mobile)"

    ua_tablet = "Mozilla/5.0 (iPad; CPU OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Version/17.0 Mobile/15E148 Safari/604.1"
    assert parse_device_from_ua(ua_tablet) == "Safari no iPad (Tablet)"


def test_auditoria_estatisticas_e_expurgo_com_backup(db):
    agora = datetime.now(timezone.utc).replace(tzinfo=None)
    data_antiga = agora - timedelta(days=400)  # > 13 meses

    # Insere registros de auditoria
    log1 = AuditLog(
        user_id=1,
        user_name="Admin",
        action="EXPORT_EXCEL",
        entity="PACIENTES",
        entity_id=None,
        device="Chrome no Windows (Desktop)",
        details=json.dumps({"modulo": "pacientes", "total_registros": 50}),
        timestamp=agora,
    )
    log2 = AuditLog(
        user_id=1,
        user_name="Admin",
        action="SEARCH",
        entity="PROCEDIMENTOS",
        entity_id=None,
        device="Safari no macOS (Desktop)",
        details=json.dumps({"termo": "raio x", "total_resultados": 5}),
        timestamp=agora,
    )
    log_velho = AuditLog(
        user_id=2,
        user_name="Operador",
        action="USER_UPDATE",
        entity="USER",
        entity_id=2,
        device="Edge no Windows (Desktop)",
        details="{}",
        timestamp=data_antiga,
    )
    db.add(log1)
    db.add(log2)
    db.add(log_velho)
    db.commit()

    # Estatísticas
    stats = obter_estatisticas_auditoria(db)
    assert stats["total_registros"] == 3
    assert stats["total_exportacoes"] == 1
    assert stats["total_pesquisas"] == 1
    assert stats["ultimos_30_dias"] == 2

    # Expurgo com retenção de 12 meses (deve expurgar log_velho de 400 dias atrás)
    resultado = executar_rotacao_logs(db, meses_retencao=12, arquivar=False)
    assert resultado["removidos"] == 1

    restantes = db.exec(select(AuditLog)).all()
    acoes = [r.action for r in restantes]
    assert "USER_UPDATE" not in acoes
    assert "EXPORT_EXCEL" in acoes
    assert "SEARCH" in acoes
    assert "AUDIT_ROTATION" in acoes


def test_user_unidade_padrao_e_atribuicao():
    from app.services.users import to_out, create_user, update_user
    from app.schemas import UserCreate, UserUpdate

    engine = create_engine("sqlite:///:memory:")
    SQLModel.metadata.create_all(engine)
    with Session(engine) as db:
        uni = UnidadeSaude(nome="PSF Aparecida do Sul", cnes="7653220", ativo=True)
        db.add(uni)
        db.commit()
        db.refresh(uni)

        data = UserCreate(
            username="operador_psf",
            password="Password123",
            gerar_temporaria=False,
            nome="Operador PSF",
            unidade_saude_id=uni.id,
            telas=["lab"],
        )
        user, _ = create_user(db, data)
        assert user.unidade_saude_id == uni.id

        out = to_out(user, db)
        assert out["unidade_saude_id"] == uni.id
        assert out["unidade_saude_nome"] == "PSF Aparecida do Sul"

        # Atualização de unidade
        uni2 = UnidadeSaude(nome="Secretaria Municipal de Saúde", cnes="6487882", ativo=True)
        db.add(uni2)
        db.commit()
        db.refresh(uni2)

        up = UserUpdate(unidade_saude_id=uni2.id)
        user_up = update_user(db, user, up)
        assert user_up.unidade_saude_id == uni2.id
        out2 = to_out(user_up, db)
        assert out2["unidade_saude_nome"] == "Secretaria Municipal de Saúde"


def test_exportar_fila_apenas_selecionados(db):
    from app.api.v1.tfd import exportar_fila

    ator = User(id=1, username="admin", password_hash="hash", is_superuser=True)
    db.add(ator)

    pac1 = Paciente(nome="Paciente Um", cpf="111.111.111-11", municipio="Guapé")
    pac2 = Paciente(nome="Paciente Dois", cpf="222.222.222-22", municipio="Guapé")
    db.add(pac1)
    db.add(pac2)
    db.commit()

    ex1 = ExameFila(
        paciente_id=pac1.id,
        paciente_nome=pac1.nome,
        exame="0201010010 · BIOPSIA",
        tipo_fila="Exames",
        prioridade="rotina",
        status="aguardando",
        ativo=True,
    )
    ex2 = ExameFila(
        paciente_id=pac2.id,
        paciente_nome=pac2.nome,
        exame="0201010020 · ULTRASSONOGRAFIA",
        tipo_fila="Exames",
        prioridade="urgente",
        status="aguardando",
        ativo=True,
    )
    db.add(ex1)
    db.add(ex2)
    db.commit()

    # Quando nenhum ID é passado, exporta todos da fila (2 pacientes)
    resp_todos = exportar_fila(db=db, ator=ator)
    assert resp_todos.status_code == 200
    import io
    import openpyxl
    wb_todos = openpyxl.load_workbook(io.BytesIO(resp_todos.body))
    ws_todos = wb_todos.active
    texto_todos = " ".join(str(c.value) for row in ws_todos.iter_rows() for c in row if c.value)
    assert "Paciente Um" in texto_todos
    assert "Paciente Dois" in texto_todos

    # Quando IDs específicos são passados, exporta apenas os selecionados
    resp_sel = exportar_fila(ids=f"{ex1.id}", db=db, ator=ator)
    assert resp_sel.status_code == 200
    wb_sel = openpyxl.load_workbook(io.BytesIO(resp_sel.body))
    ws_sel = wb_sel.active
    texto_sel = " ".join(str(c.value) for row in ws_sel.iter_rows() for c in row if c.value)
    assert "Paciente Um" in texto_sel
    assert "Paciente Dois" not in texto_sel


def test_exportar_fila_psf_apenas_selecionados(db):
    from app.api.v1.lab_fase6 import exportar_fila_psf
    from app.models.lab_fase6 import FilaPsf

    ator = User(id=1, username="admin", password_hash="hash", is_superuser=True)
    db.add(ator)

    pac1 = Paciente(nome="Paciente Lab 1", cpf="333.333.333-33", municipio="Guapé")
    pac2 = Paciente(nome="Paciente Lab 2", cpf="444.444.444-44", municipio="Guapé")
    db.add(pac1)
    db.add(pac2)
    db.commit()

    f1 = FilaPsf(
        paciente_id=pac1.id,
        unidade_psf="PSF Central",
        medico_solicitante="Médico da Unidade",
        status="aguardando",
    )
    f2 = FilaPsf(
        paciente_id=pac2.id,
        unidade_psf="PSF Bela Vista",
        medico_solicitante="Médico da Unidade",
        status="aguardando",
    )
    db.add(f1)
    db.add(f2)
    db.commit()

    # Exporta apenas o f2
    resp = exportar_fila_psf(ids=f"{f2.id}", db=db, ator=ator)
    assert resp.status_code == 200
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.body))
    ws = wb.active
    texto = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Paciente Lab 2" in texto
    assert "Paciente Lab 1" not in texto


# =====================================================================
# 6. REGRAS DE SOLICITAÇÃO (PSF vs. TFD), RECEPÇÃO EM LOTE E RELATÓRIO
# =====================================================================

def test_solicitacao_regras_psf_restricao_unidade_e_data(db):
    """Operadores de PSF:
    - Só podem lançar para sua própria unidade de saúde vinculada.
    - Só podem lançar para a data atual (date.today()).
    """
    uni1 = UnidadeSaude(nome="PSF Centro I", cnes="1111111", ativo=True)
    uni2 = UnidadeSaude(nome="PSF Bela Vista II", cnes="2222222", ativo=True)
    pac = Paciente(nome="Carlos Eduardo", cpf="777.777.777-77", municipio="Guapé")
    db.add_all([uni1, uni2, pac])
    db.commit()

    operador_psf = User(
        username="op_psf1",
        password_hash="hash",
        telas="tfd_fila_lancamento",
        unidade_saude_id=uni1.id,
        is_superuser=False,
    )
    db.add(operador_psf)
    db.commit()

    # 1. Tentativa de lançar em data retroativa por operador de PSF deve gerar 403
    ex_data_invalida = ExameIn(
        paciente_id=pac.id,
        exame="ELETROCARDIOGRAMA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        justificativa="Acompanhamento ambulatorial de rotina",
        data_pedido=date.today() - timedelta(days=2),
        unidade_saude_id=uni1.id,
    )
    with pytest.raises(HTTPException) as exc_info:
        criar_exame(ex_data_invalida, db=db, ator=operador_psf)
    assert exc_info.value.status_code == 403
    assert "dia atual" in exc_info.value.detail

    # 2. Tentativa de lançar para outra unidade de saúde deve gerar 403
    ex_unidade_invalida = ExameIn(
        paciente_id=pac.id,
        exame="ELETROCARDIOGRAMA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        justificativa="Acompanhamento ambulatorial de rotina",
        data_pedido=date.today(),
        unidade_saude_id=uni2.id,
    )
    with pytest.raises(HTTPException) as exc_info:
        criar_exame(ex_unidade_invalida, db=db, ator=operador_psf)
    assert exc_info.value.status_code == 403
    assert "própria unidade" in exc_info.value.detail

    # 3. Lançamento correto na data de hoje e na sua própria unidade tem sucesso
    ex_valido = ExameIn(
        paciente_id=pac.id,
        exame="ELETROCARDIOGRAMA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        justificativa="Acompanhamento ambulatorial de rotina",
        data_pedido=date.today(),
        unidade_saude_id=uni1.id,
    )
    salvo = criar_exame(ex_valido, db=db, ator=operador_psf)
    assert salvo.id is not None
    assert salvo.unidade_saude_id == uni1.id
    assert salvo.unidade_nome == "PSF Centro I"
    assert salvo.unidade_cnes == "1111111"


def test_solicitacao_tfd_central_permissao_irrestrita(db):
    """Operador TFD Central / Regulador:
    - Pode lançar com data livre.
    - Pode selecionar qualquer unidade de saúde.
    """
    uni = UnidadeSaude(nome="PSF Pontal", cnes="3333333", ativo=True)
    pac = Paciente(nome="Mariana Lima", cpf="888.888.888-88", municipio="Guapé")
    db.add_all([uni, pac])
    db.commit()

    operador_tfd = User(
        username="op_tfd_central",
        password_hash="hash",
        telas="tfd,tfd_exames,tfd_agenda",
        unidade_saude_id=None,
        is_superuser=False,
    )
    db.add(operador_tfd)
    db.commit()

    ex_tfd = ExameIn(
        paciente_id=pac.id,
        exame="CONSULTA OFTALMOLOGICA",
        tipo_fila="Oftalmologia",
        prioridade="urgente",
        justificativa="Perda súbita de acuidade visual esquerda",
        data_pedido=date.today() - timedelta(days=7),
        unidade_saude_id=uni.id,
    )
    salvo = criar_exame(ex_tfd, db=db, ator=operador_tfd)
    assert salvo.id is not None
    assert salvo.data_pedido == date.today() - timedelta(days=7)
    assert salvo.unidade_saude_id == uni.id
    assert salvo.unidade_nome == "PSF Pontal"


def test_filtro_unidade_fila_e_agendamentos(db):
    """Filtro por unidade na fila de espera."""
    uni_a = UnidadeSaude(nome="PSF Alpha", cnes="1010101", ativo=True)
    uni_b = UnidadeSaude(nome="PSF Beta", cnes="2020202", ativo=True)
    pac1 = Paciente(nome="Paciente Alpha", cpf="101.101.101-10", municipio="Guapé")
    pac2 = Paciente(nome="Paciente Beta", cpf="202.202.202-20", municipio="Guapé")
    db.add_all([uni_a, uni_b, pac1, pac2])
    db.commit()

    ex_a = ExameFila(
        paciente_id=pac1.id,
        paciente_nome=pac1.nome,
        exame="RAIO-X DE TORAX",
        tipo_fila="Raio-X",
        prioridade="rotina",
        status="aguardando",
        data_pedido=date.today(),
        unidade_saude_id=uni_a.id,
        unidade_nome=uni_a.nome,
        ativo=True,
    )
    ex_b = ExameFila(
        paciente_id=pac2.id,
        paciente_nome=pac2.nome,
        exame="TOMOGRAFIA COMPUTADORIZADA",
        tipo_fila="Tomografia",
        prioridade="rotina",
        status="aguardando",
        data_pedido=date.today(),
        unidade_saude_id=uni_b.id,
        unidade_nome=uni_b.nome,
        ativo=True,
    )
    db.add_all([ex_a, ex_b])
    db.commit()

    admin = User(username="admin_filtro", password_hash="hash", is_superuser=True)

    res_a = listar_exames(unidade_id=uni_a.id, db=db, ator=admin)
    itens_a = res_a["data"] if isinstance(res_a, dict) else res_a
    assert any(e.id == ex_a.id for e in itens_a)
    assert not any(e.id == ex_b.id for e in itens_a)

    res_b = listar_exames(unidade_id=uni_b.id, db=db, ator=admin)
    itens_b = res_b["data"] if isinstance(res_b, dict) else res_b
    assert any(e.id == ex_b.id for e in itens_b)
    assert not any(e.id == ex_a.id for e in itens_b)


def test_recepcao_em_lote_e_status_atendimento(db):
    """Recepção em lote e atualização do desfecho do atendimento."""
    from app.api.v1.tfd import (
        recepcionar_exames_lote,
        definir_status_atendimento_lote,
        ExameIdsIn,
        AtendimentoStatusLoteIn,
    )

    pac = Paciente(nome="Juliana Paes", cpf="303.303.303-30", municipio="Guapé")
    db.add(pac)
    db.commit()

    ex1 = ExameFila(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="ULTRASSONOGRAFIA ABDOMINAL",
        tipo_fila="Ultrassom",
        prioridade="rotina",
        status="agendado",
        data_agendada=date.today(),
        hora_agendada="08:00",
        ativo=True,
    )
    ex2 = ExameFila(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="CONSULTA CARDIOLOGIA",
        tipo_fila="Cardiologia",
        prioridade="rotina",
        status="agendado",
        data_agendada=date.today(),
        hora_agendada="08:30",
        ativo=True,
    )
    db.add_all([ex1, ex2])
    db.commit()

    operador = User(username="recepcao", password_hash="hash", telas="tfd_agenda", is_superuser=True)

    # 1. Validação de CID obrigatório: deve barrar se não houver CID
    import pytest
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc_info:
        recepcionar_exames_lote(ExameIdsIn(ids=[ex1.id, ex2.id]), db=db, ator=operador)
    assert exc_info.value.status_code == 400
    assert "exige(m) o preenchimento de CID" in exc_info.value.detail

    # 2. Recepção em lote com preenchimento de CID
    resp_lote = recepcionar_exames_lote(
        ExameIdsIn(ids=[ex1.id, ex2.id], cids={str(ex1.id): "Z00", str(ex2.id): "I10"}),
        db=db,
        ator=operador,
    )
    assert resp_lote["ok"] is True
    assert resp_lote["atualizados"] == 2

    db.refresh(ex1)
    db.refresh(ex2)
    assert ex1.status == "recepcionado"
    assert ex1.recepcionado_em is not None
    assert ex2.status == "recepcionado"

    # 2. Atualizar status de atendimento para realizado
    resp_atend1 = definir_status_atendimento_lote(
        AtendimentoStatusLoteIn(ids=[ex1.id], status="realizado", resultado="Exame concluído com laudo"),
        db=db,
        ator=operador,
    )
    assert resp_atend1["ok"] is True
    db.refresh(ex1)
    assert ex1.status == "realizado"
    assert "concluído com laudo" in ex1.resultado

    # 3. Atualizar status de atendimento para faltou
    resp_atend2 = definir_status_atendimento_lote(
        AtendimentoStatusLoteIn(ids=[ex2.id], status="faltou"),
        db=db,
        ator=operador,
    )
    assert resp_atend2["ok"] is True
    db.refresh(ex2)
    assert ex2.status == "faltou"


def test_relatorio_agendados_metricas_e_desfechos(db):
    """Relatório gerencial de agendados com cálculo de absenteísmo e exportação."""
    from app.api.v1.tfd import obter_relatorio_agendados

    pac = Paciente(nome="Roberto Carlos", cpf="505.505.505-50", municipio="Guapé")
    db.add(pac)
    db.commit()

    hoje = date.today()
    e_realizado = ExameFila(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="ENDOSCOPIA DIGESTIVA",
        tipo_fila="Endoscopia",
        prioridade="rotina",
        status="realizado",
        data_agendada=hoje,
        hora_agendada="09:00",
        ativo=True,
    )
    e_faltou = ExameFila(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="COLONOSCOPIA",
        tipo_fila="Endoscopia",
        prioridade="rotina",
        status="faltou",
        data_agendada=hoje,
        hora_agendada="10:00",
        ativo=True,
    )
    e_agendado = ExameFila(
        paciente_id=pac.id,
        paciente_nome=pac.nome,
        exame="CONSULTA GASTRO",
        tipo_fila="Consultas",
        prioridade="rotina",
        status="agendado",
        data_agendada=hoje,
        hora_agendada="11:00",
        ativo=True,
    )
    db.add_all([e_realizado, e_faltou, e_agendado])
    db.commit()

    ator = User(username="gestor", password_hash="hash", telas="tfd_agenda", is_superuser=True)
    rel = obter_relatorio_agendados(
        data_de=hoje.isoformat(),
        data_ate=hoje.isoformat(),
        db=db,
        ator=ator,
    )

    metricas = rel["metricas"]
    assert metricas["total"] >= 3
    assert metricas["realizados"] >= 1
    assert metricas["faltas"] >= 1
    assert metricas["agendados"] >= 1
    # Verifica que absenteísmo e comparecimento são calculados
    assert "absenteismo_pct" in metricas
    assert "comparecimento_pct" in metricas
    assert len(rel["itens"]) >= 3



