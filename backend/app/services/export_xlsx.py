"""Exportação XLSX padronizada (openpyxl).

Cabeçalho em todo relatório: identificação, data/hora da geração e
usuário solicitante. Colunas visíveis com formatação (datas DD/MM/AAAA,
CPF/códigos como texto p/ preservar zeros e pontuação).
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_LINHAS = 20000

TITULO_FONT = Font(bold=True, size=14, color="0A3D62")
META_FONT = Font(size=10, color="555555")
HEADER_FILL = PatternFill("solid", fgColor="0A3D62")
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def fmt_data_br(valor) -> str:
    if valor is None or valor == "":
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    s = str(valor).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
    return s


def gerar_xlsx(
    titulo: str,
    usuario_nome: str,
    colunas: list[dict],
    linhas: list[dict],
) -> bytes:
    """colunas: [{chave, rotulo, largura?, formato?: texto|data|numero|moeda}]."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"[:31]
    gerado = datetime.now()

    ws["A1"] = titulo
    ws["A1"].font = TITULO_FONT
    ws["A2"] = f"Gerado em: {gerado.strftime('%d/%m/%Y %H:%M')}   ·   Solicitado por: {usuario_nome or '—'}"
    ws["A2"].font = META_FONT
    ws["A3"] = f"Registros: {len(linhas)}"
    ws["A3"].font = META_FONT

    base = 5
    for i, c in enumerate(colunas, start=1):
        cell = ws.cell(row=base, column=i, value=c.get("rotulo", c.get("chave", "")))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(i)].width = c.get("largura", 20)
    ws.row_dimensions[base].height = 30

    for li, lin in enumerate(linhas[:MAX_LINHAS]):
        r = base + 1 + li
        zebra = li % 2 == 1
        for ci, c in enumerate(colunas, start=1):
            v = lin.get(c["chave"], "")
            formato = c.get("formato", "texto")
            cell = ws.cell(row=r, column=ci)
            if formato == "data":
                cell.value = fmt_data_br(v)
                cell.number_format = "@"
            elif formato == "numero":
                try:
                    cell.value = float(v) if v not in (None, "") else None
                    cell.number_format = "#,##0.00"
                except (TypeError, ValueError):
                    cell.value = "" if v is None else str(v)
            elif formato == "moeda":
                try:
                    cell.value = float(v) if v not in (None, "") else None
                    cell.number_format = 'R$ #,##0.00'
                except (TypeError, ValueError):
                    cell.value = "" if v is None else str(v)
            else:  # texto: CPF, códigos e textos preservados
                cell.value = "" if v is None else str(v)
                cell.number_format = "@"
            cell.font = Font(size=10)
            cell.border = BORDER_THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if zebra:
                cell.fill = PatternFill("solid", fgColor="F5F7F8")
    ultima_col = get_column_letter(len(colunas))
    ws.auto_filter.ref = f"A{base}:{ultima_col}{base + min(len(linhas), MAX_LINHAS)}"
    ws.freeze_panes = f"A{base + 1}"
    ws.sheet_properties.pageSetUpPr = None
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nome_arquivo(base: str) -> str:
    return f"{base}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"


def resposta_xlsx(titulo: str, usuario_nome: str, colunas: list[dict], linhas: list[dict], arquivo_base: str):
    from fastapi import Response

    return Response(
        content=gerar_xlsx(titulo, usuario_nome, colunas, linhas),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{nome_arquivo(arquivo_base)}"',
            "Cache-Control": "no-store",
        },
    )
