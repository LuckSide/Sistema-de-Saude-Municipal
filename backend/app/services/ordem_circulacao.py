"""Ordem de Circulação de Veículos (.docx) — réplica do modelo oficial.

Estrutura idêntica a "Viagens padrão.doc": título, cabeçalho, dados da
viagem, DESPESAS DE VIAGEM + vistos/assinaturas e RELAÇÃO DE PASSAGEIROS.
Única diferença autorizada: 3ª coluna da tabela de passageiros
(Pront. ➔ Data de Nasc.).
"""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _borda_fina(tab) -> None:
    """Grade fina preta (1px) em todas as bordas — separação visual das colunas."""
    from docx.oxml import OxmlElement

    tblPr = tab._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        bordas.append(el)
    tblPr.append(bordas)


def _larguras(tab, larguras_cm: list[float]) -> None:
    tab.autofit = False
    tab.allow_autofit = False
    for i, larg in enumerate(larguras_cm):
        for cel in tab.columns[i].cells:
            cel.width = Cm(larg)


def _altura_minima(linha, cm: float = 0.6) -> None:
    from docx.oxml import OxmlElement

    trPr = linha._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def _celula(linha, idx, texto, negrito=False, centralizar=False, tamanho=9):
    cel = linha.cells[idx]
    cel.text = ""
    p = cel.paragraphs[0]
    if centralizar:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto or "")
    run.font.size = Pt(tamanho)
    run.bold = negrito
    return cel


def _tabela(doc, cabecalhos, linhas, larguras=None, altura_linhas=None):
    tab = doc.add_table(rows=1 + len(linhas), cols=len(cabecalhos))
    tab.style = "Table Grid"
    _borda_fina(tab)
    if larguras:
        _larguras(tab, larguras)
    for i, h in enumerate(cabecalhos):
        _celula(tab.rows[0], i, h, negrito=True, centralizar=True)
    for r, lin in enumerate(linhas, start=1):
        for i, v in enumerate(lin):
            _celula(tab.rows[r], i, v, centralizar=(i == 0))
        if altura_linhas:
            _altura_minima(tab.rows[r], altura_linhas)
    return tab


def _titulo(doc, texto, tamanho=13):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(tamanho)
    return p


def _data_br(iso: str) -> str:
    s = (iso or "")[:10]
    return f"{s[8:10]}/{s[5:7]}/{s[0:4]}" if len(s) == 10 and s[4] == "-" else (iso or "")


def _nasc_br(iso: str) -> str:
    s = (iso or "")[:10]
    if len(s) == 10 and s[4] == "-":
        return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
    return (iso or "") or "—"


def gerar_ordem(
    veiculo_modelo: str,
    placa: str,
    destino: str,
    motorista: str,
    departamento: str,
    data_saida: str,
    hora_saida: str,
    data_retorno: str,
    hora_retorno: str,
    km_inicial: int,
    km_final: int,
    adiantamento: float,
    passageiros: list[dict],
    horas_viagem: str = "",
) -> bytes:
    doc = Document()
    for secao in doc.sections:
        secao.top_margin = Pt(36)
        secao.bottom_margin = Pt(36)
        secao.left_margin = Cm(1.5)
        secao.right_margin = Cm(1.5)

    _titulo(doc, "ORDEM DE CIRCULAÇÃO DE VEÍCULOS E PRESTAÇÃO DE CONTAS DE VIAGEM")

    doc.add_paragraph()
    _tabela(doc, ["Campo", "Valor"], [
        ["VEICULO", veiculo_modelo or "—"],
        ["PLACA", placa or "—"],
        ["DESTINO", destino or "—"],
        ["MOTORISTA/FAVORECIDO", motorista or "—"],
        ["DEPARTAMENTO", departamento or "—"],
    ])

    doc.add_paragraph()
    km_rodados = (km_final - km_inicial) if km_final and km_inicial and km_final >= km_inicial else ""
    _tabela(doc, ["Campo", "Valor"], [
        ["Data saída", _data_br(data_saida) or "—"],
        ["Hora saída", hora_saida or "—"],
        ["Km inicial", str(km_inicial or "—")],
        ["Data retorno", _data_br(data_retorno) or "—"],
        ["Hora retorno", hora_retorno or "—"],
        ["Km final", str(km_final or "—")],
        ["Adiantamento", f"{adiantamento:.2f}".replace(".", ",") if adiantamento else "—"],
        ["Horas viagem", horas_viagem or "—"],
        ["Km rodados", str(km_rodados) if km_rodados != "" else "—"],
    ])

    doc.add_paragraph()
    _titulo(doc, "DESPESAS DE VIAGEM", tamanho=11)
    _tabela(doc, ["Despesa", "Valor (R$)", "Vistos"], [
        ["Alimentação", "", ""],
        ["Hospedagem", "", ""],
        ["Combustível", "", ""],
        ["Peças", "", ""],
        ["Serviços", "", ""],
        ["Táxi", "", ""],
        ["Ônibus", "", ""],
        ["Pedágio", "", ""],
        ["Outros", "", ""],
        ["TOTAL", "", ""],
    ])

    doc.add_paragraph()
    for rotulo in ["Favorecido", "Chefe do Departamento", "Controle Interno", "Assessoria Jurídica"]:
        pa = doc.add_paragraph()
        pa.add_run(f"{rotulo}: ").bold = True
        pa.add_run("_" * 50)

    # Bloco isolado: quebra de página antes da relação (modelo tem 2 páginas).
    doc.add_page_break()
    _titulo(doc, "RELAÇÃO DE PASSAGEIROS", tamanho=11)
    linhas = []
    for i, ps in enumerate(passageiros, start=1):
        linhas.append([
            f"{i:02d}",
            ps.get("paciente_nome", "") or "—",
            _nasc_br(ps.get("paciente_nascimento", "")),
            ps.get("tratamento", "") or "—",
            ps.get("unidade", "") or "—",
            ps.get("hora_consulta", "") or "—",
            destino or "—",
        ])
    # 15 linhas fixas p/ preenchimento manual limpo (completa com vazias).
    while len(linhas) < 15:
        linhas.append([f"{len(linhas) + 1:02d}", "", "", "", "", "", ""])
    _tabela(doc, ["N", "Passageiro", "Data nasc.", "Tratamento", "Unidade",
                  "Horário consulta", "Cidade"], linhas,
            larguras=[0.9, 3.6, 2.0, 2.8, 2.8, 1.8, 2.6], altura_linhas=0.6)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
